"""API：用 tests/fakes.py 的替身，不需金鑰。只驗證流程與錯誤碼，不驗證規則效果。"""

import json
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from coach.api import create_app
from coach.flow import CoachService
from coach.inputs import parse_industries
from coach.rules import load_rules
from coach.store import JsonFileStore
from tests.fakes import SequenceLLM, clarify, vfp

BODY = {"name": "王主管", "industry": 9}
ITEMS = [{"task": "a", "detail": "b", "product": "p"}, {"task": "c", "detail": "d", "product": "p"}]
ITEMS_OUT = [{**i, "assignee": None} for i in ITEMS]  # 讀取時會補上 assignee 預設值


@pytest.fixture
def make(tmp_path):
    def _make(*replies):
        llm = SequenceLLM(*replies)
        svc = CoachService(JsonFileStore(tmp_path / "sessions"), lambda: llm)
        return TestClient(create_app(svc)), llm

    return _make


def test_industries_come_from_rules_file(make):
    client, _ = make()
    r = client.get("/industries")
    assert r.status_code == 200
    items = parse_industries(load_rules().text)
    assert r.json() == [{"no": i, "name": n} for i, n in enumerate(items, start=1)]
    assert len(r.json()) == 20


def test_only_name_and_industry_are_required(make):
    client, llm = make(clarify("Q"))
    r = client.post("/sessions", json=BODY)
    assert r.status_code == 200 and r.json()["status"] == "clarify"
    assert llm.calls[0][1]["content"].count("（未填）") == 2


@pytest.mark.parametrize(
    "body,msg",
    [
        ({"industry": 9}, "name缺少"),
        ({"name": "王主管"}, "industry缺少"),
        ({"name": "王主管", "industry": "九"}, "industry格式不正確"),
        ({"name": "  ", "industry": 9}, "稱呼不可空白"),
        ({"name": "王主管", "industry": 0}, "產業必須是 1～20 的編號"),
        ({"name": "王主管", "industry": 21}, "產業必須是 1～20 的編號"),
        ({"name": "王主管", "industry": 9, "score": 11}, "營運自評分數必須是 1～10 的整數"),
    ],
)
def test_start_validation_returns_422_in_chinese_without_calling_model(make, body, msg):
    client, llm = make(clarify("Q"))
    r = client.post("/sessions", json=body)
    assert r.status_code == 422 and r.json() == {"detail": msg}
    assert llm.calls == []


def test_industry_20_requires_text_and_is_stored_with_prefix(make):
    client, llm = make(clarify("Q"))
    r = client.post("/sessions", json={"name": "王主管", "industry": 20})
    assert r.status_code == 422 and r.json() == {"detail": "選「其他」時必須填寫產業"}
    r = client.post("/sessions", json={"name": "王主管", "industry": 20, "industry_other": " 修繕 "})
    assert r.status_code == 200
    assert client.get(f"/sessions/{r.json()['session_id']}").json()["inputs"]["industry"] == "其他：修繕"
    assert "其他：修繕" in llm.calls[0][1]["content"]


def test_clarify_answer_vfp_flow(make):
    client, llm = make(clarify("Q1"), clarify("Q2"), vfp("V", ITEMS))
    sid = client.post("/sessions", json={**BODY, "product": "P", "score": 6, "reason": "R"}).json()["session_id"]
    r = client.post(f"/sessions/{sid}/answer", json={"answer": "A1"})
    assert r.json()["status"] == "clarify" and r.json()["question"] == "Q2"
    r = client.post(f"/sessions/{sid}/answer", json={"answer": "A2"})
    body = r.json()
    assert (body["status"], body["vfp"], body["byproducts"], body["edited"]) == ("vfp", "V", ITEMS_OUT, False)
    assert body["rules"]["version"] == load_rules().version
    s = client.get(f"/sessions/{sid}").json()
    assert s["status"] == "vfp" and s["inputs"]["score"] == 6 and s["inputs"]["reason"] == "R"
    assert [e["manager"] for e in s["exchanges"]] == ["A1", "A2"]
    assert len(llm.calls) == 3


def test_first_round_can_be_vfp(make):
    client, _ = make(vfp("V", ITEMS))
    assert client.post("/sessions", json=BODY).json()["status"] == "vfp"


def test_answer_validation_and_unknown_session(make):
    client, _ = make(clarify("Q"))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    assert client.post(f"/sessions/{sid}/answer", json={}).json() == {"detail": "answer缺少"}
    assert client.post(f"/sessions/{sid}/answer", json={"answer": "  "}).json() == {"detail": "回答不可為空"}
    for call in (
        lambda: client.get("/sessions/nope"),
        lambda: client.post("/sessions/nope/answer", json={"answer": "x"}),
        lambda: client.put("/sessions/nope/byproducts", json={"vfp": "V", "byproducts": ITEMS}),
    ):
        r = call()
        assert r.status_code == 404 and "找不到" in r.json()["detail"]


def test_answer_after_vfp_is_409(make):
    client, llm = make(vfp("V", ITEMS))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    r = client.post(f"/sessions/{sid}/answer", json={"answer": "more"})
    assert r.status_code == 409 and "沒有待回答的問題" in r.json()["detail"]
    assert len(llm.calls) == 1


def test_put_byproducts_is_409_unless_vfp(make):
    client, _ = make(clarify("Q"))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    r = client.put(f"/sessions/{sid}/byproducts", json={"vfp": "V", "byproducts": ITEMS})
    assert r.status_code == 409 and "只有 vfp 狀態能修改" in r.json()["detail"]


def test_put_overwrites_and_get_shows_original_and_edited(make):
    client, _ = make(vfp("原始", ITEMS))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    new = [{"task": "x", "detail": "y", "product": "p"}]
    r = client.put(f"/sessions/{sid}/byproducts", json={"vfp": "改過", "byproducts": new})
    assert r.status_code == 200
    assert (r.json()["status"], r.json()["vfp"], r.json()["byproducts"], r.json()["edited"]) == ("vfp", "改過", [{**n, "assignee": None} for n in new], True)
    result = client.get(f"/sessions/{sid}").json()["result"]
    assert result["original"] == {"vfp": "原始", "byproducts": ITEMS_OUT}
    assert result["edited"]["vfp"] == "改過" and result["edited"]["byproducts"] == [{**n, "assignee": None} for n in new]
    # 再改一次，原始版本仍在；之後 answer 仍是 409
    client.put(f"/sessions/{sid}/byproducts", json={"vfp": "又改", "byproducts": ITEMS})
    result = client.get(f"/sessions/{sid}").json()["result"]
    assert result["original"]["vfp"] == "原始" and result["edited"]["vfp"] == "又改"
    assert client.post(f"/sessions/{sid}/answer", json={"answer": "x"}).status_code == 409


@pytest.mark.parametrize(
    "body",
    [
        {"vfp": "", "byproducts": ITEMS},
        {"vfp": "V", "byproducts": []},
        {"vfp": "V", "byproducts": [{"task": "a", "detail": "b", "product": ""}]},
        {"vfp": "V", "byproducts": [{"task": " ", "detail": "b", "product": "p"}]},
    ],
)
def test_put_validation(make, body):
    client, _ = make(vfp("V", ITEMS))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    r = client.put(f"/sessions/{sid}/byproducts", json=body)
    assert r.status_code == 422
    assert client.get(f"/sessions/{sid}").json()["result"]["edited"] is None


def test_put_missing_fields(make):
    client, _ = make(vfp("V", ITEMS))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    r = client.put(f"/sessions/{sid}/byproducts", json={"vfp": "V", "byproducts": [{"task": "a", "detail": "b"}, {"task": "c", "detail": "d", "product": " "}]})
    assert r.status_code == 422 and r.json()["detail"] == "第 1 列的副產品不可為空"
    r = client.put(f"/sessions/{sid}/byproducts", json={"vfp": "V", "byproducts": [{"task": "a", "product": "p"}, {"detail": "d", "product": "p"}]})
    assert r.status_code == 422 and r.json()["detail"] == "第 2 列的工作項目不可為空"


@pytest.mark.parametrize("bad", ["not json", '{"status":"vfp","vfp":"V"}', '{"status":"clarify"}'])
def test_unparseable_model_reply_is_502_not_saved_and_resendable(make, tmp_path, bad):
    client, _ = make(bad, clarify("Q"))
    r = client.post("/sessions", json=BODY)
    assert r.status_code == 502 and "無法解析" in r.json()["detail"]
    assert not list((tmp_path / "sessions").glob("*.json"))
    assert client.post("/sessions", json=BODY).status_code == 200  # 同一請求重送


def test_unparseable_reply_on_answer_keeps_session_and_can_resend(make):
    client, _ = make(clarify("Q1"), "not json", vfp("V", ITEMS))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    before = client.get(f"/sessions/{sid}").json()
    assert client.post(f"/sessions/{sid}/answer", json={"answer": "A1"}).status_code == 502
    assert client.get(f"/sessions/{sid}").json() == before
    assert client.post(f"/sessions/{sid}/answer", json={"answer": "A1"}).json()["status"] == "vfp"


def test_model_unavailable_is_503(tmp_path):
    def boom():
        raise RuntimeError("缺少環境變數")

    client = TestClient(create_app(CoachService(JsonFileStore(tmp_path / "sessions"), boom)))
    r = client.post("/sessions", json=BODY)
    assert r.status_code == 503 and "缺少環境變數" in r.json()["detail"]


def test_docs_page_and_endpoints_listed(make):
    client, _ = make()
    assert client.get("/docs").status_code == 200
    paths = client.get("/openapi.json").json()["paths"]
    assert set(paths) == {"/industries", "/sessions", "/sessions/{session_id}/answer", "/sessions/{session_id}/byproducts", "/sessions/{session_id}/staff", "/sessions/{session_id}", "/sessions/{session_id}/export.xlsx"}


# ---- 網頁 -----------------------------------------------------------------


def test_index_page_is_served_with_title_and_local_assets(make):
    client, _ = make()
    r = client.get("/")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/html")
    assert "管到飛" in r.text and "每週陪跑教練，一上任就帶大家起飛" in r.text
    assert "/static/style.css" in r.text and "/static/app.js" in r.text


def test_index_page_uses_no_external_resources(make):
    import re

    client, _ = make()
    for path in ("/", "/static/style.css", "/static/app.js"):
        text = client.get(path).text
        assert not re.search(r"(https?:)?//[\w.-]+\.[a-z]{2,}", re.sub(r"http://127\.0\.0\.1", "", text)), path


@pytest.mark.parametrize("path,kind", [("/static/style.css", "text/css"), ("/static/app.js", "javascript")])
def test_static_files_are_available(make, path, kind):
    client, _ = make()
    r = client.get(path)
    assert r.status_code == 200 and kind in r.headers["content-type"]


def test_css_defines_the_palette_variables_and_no_pure_black(make):
    client, _ = make()
    css = client.get("/static/style.css").text
    for name, value in {
        "--navy-900": "#0C3467", "--navy-800": "#0F325E", "--blue-700": "#1E5AA8", "--slate-700": "#2D4A6E",
        "--gray-600": "#5B6B80", "--gray-blue-400": "#8A9BB0", "--sky-300": "#9CC3EA", "--sky-100": "#DCEBFA",
        "--mist-100": "#E8EEF5", "--yellow-400": "#F4C84D",
    }.items():
        assert f"{name}: {value};" in css
    assert "#000" not in css.lower().replace("#0c3467", "").replace("#0f325e", "")


def test_missing_logo_is_just_404_and_page_still_loads(make):
    client, _ = make()
    assert client.get("/static/logo.png").status_code in (200, 404)
    assert client.get("/").status_code == 200


def test_api_routes_still_work_next_to_static_mount(make):
    client, _ = make()
    assert client.get("/industries").status_code == 200
    assert client.get("/docs").status_code == 200


# ---- 三頁網頁與匯出 --------------------------------------------------------

WEB = Path(__file__).resolve().parent.parent / "coach" / "web"


def test_page_mentions_three_services_and_routes(make):
    client, _ = make()
    html = client.get("/").text
    js = client.get("/static/app.js").text
    for name in ("職位的帽子", "產品鑑定", "行政管理等級表"):
        assert name in html or name in js
    for route in ("#/chat", "#/result"):
        assert route in html or route in js
    assert 'aria-disabled="true"' in html and html.count("即將推出") == 2


def test_chat_questions_match_rules_chapter_2_wording(make):
    client, _ = make()
    js = client.get("/static/app.js").text
    for q in ("該如何稱呼您？", "你在哪個產業？", "你覺得自己/團隊/公司在賣什麼產品？", "賣的如何？", "為自己的營運打個分數（1~10分）？", "為何給這個分數？"):
        assert q in js


def test_web_files_have_no_external_urls():
    import re

    files = [p for p in WEB.rglob("*") if p.is_file() and p.suffix in {".html", ".css", ".js", ".svg"}]
    assert files
    for p in files:
        text = p.read_text(encoding="utf-8")
        urls = [u for u in re.findall(r"https?://[^\s\"')<>]+", text) if not u.startswith("http://www.w3.org/")]  # SVG xmlns 除外
        assert not urls, (p.name, urls)
        assert not re.search(r"(src|href)=[\"']//", text), p.name
        assert "@import" not in text, p.name


def vfp_session(client, edited=False):
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    if edited:
        client.put(f"/sessions/{sid}/byproducts", json={"vfp": "改過的VFP", "byproducts": [{"task": "新甲", "detail": "新說明甲", "product": "p"}, {"task": "新乙", "detail": "新說明乙", "product": "p"}, {"task": "新丙", "detail": "新說明丙", "product": "p"}]})
    return sid


def read_xlsx(response):
    import io

    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(response.content))


def test_export_xlsx_matches_model_version_when_not_edited(make):
    client, _ = make(vfp("原始VFP", ITEMS))
    sid = vfp_session(client)
    r = client.get(f"/sessions/{sid}/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ws = read_xlsx(r)["產品鑑定"]
    assert (ws["A2"].value, ws["B2"].value) == ("VFP", "原始VFP")
    assert [c.value for c in ws[3]] == ["序號", "工作項目", "工作內容", "副產品", "負責人"]
    assert [[c.value for c in ws[i]] for i in (4, 5)] == [[1, "a", "b", "p", "（未指派）"], [2, "c", "d", "p", "（未指派）"]]
    assert "由管到飛產生，日期：" in ws["A7"].value
    assert ws["A3"].fill.fgColor.rgb.endswith("0F325E") and ws["A3"].font.color.rgb.endswith("FFFFFF")
    # VFP 列是金色、navy 粗體；表頭仍是 navy-800 底白字
    assert ws["A2"].fill.fgColor.rgb.endswith("F4C84D") and ws["B2"].fill.fgColor.rgb.endswith("FBE08A")
    assert ws["A2"].font.bold and ws["A2"].font.color.rgb.endswith("0C3467") and ws["B2"].font.color.rgb.endswith("0C3467")


def test_export_xlsx_uses_edited_version(make):
    client, _ = make(vfp("原始VFP", ITEMS))
    sid = vfp_session(client, edited=True)
    ws = read_xlsx(client.get(f"/sessions/{sid}/export.xlsx"))["產品鑑定"]
    assert ws["B2"].value == "改過的VFP"
    assert [ws.cell(row=i, column=2).value for i in (4, 5, 6)] == ["新甲", "新乙", "新丙"]
    assert [ws.cell(row=i, column=1).value for i in (4, 5, 6)] == [1, 2, 3]


def test_export_filename_header_handles_chinese(make):
    from urllib.parse import quote

    client, _ = make(vfp("V", ITEMS))
    sid = vfp_session(client)
    cd = client.get(f"/sessions/{sid}/export.xlsx").headers["content-disposition"]
    assert cd.startswith("attachment;") and "filename*=UTF-8''" in cd
    assert quote("管到飛_產品鑑定_") in cd and cd.endswith(".xlsx")


def test_export_errors(make):
    client, _ = make(clarify("Q"))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    r = client.get(f"/sessions/{sid}/export.xlsx")
    assert r.status_code == 409 and "還沒有可匯出的結果" in r.json()["detail"]
    r = client.get("/sessions/nope/export.xlsx")
    assert r.status_code == 404


def test_icons_exist_and_are_valid_xml():
    import xml.etree.ElementTree as ET

    for name in ("hat", "gold", "wand"):
        path = WEB / "assets" / "icons" / f"{name}.svg"
        assert path.is_file(), name
        root = ET.parse(path).getroot()
        assert root.tag.endswith("svg") and root.get("viewBox") == "0 0 96 96"


def test_home_page_references_icons_as_img_files(make):
    client, _ = make()
    html = client.get("/").text
    for name in ("hat", "gold", "wand"):
        assert f'src="/static/assets/icons/{name}.svg"' in html
        assert client.get(f"/static/assets/icons/{name}.svg").status_code == 200
    assert html.count("<svg") == 2  # 只有 VFP 標籤旁的小星形；三個服務圖示都是 <img> 檔案
    assert "baton" not in html


def test_only_appraisal_card_is_clickable(make):
    client, _ = make()
    html = client.get("/").text
    assert html.count('class="service active"') == 1 and html.count('class="service disabled"') == 2
    assert 'href="#/chat"' in html


def test_image_export_script_is_served_and_wired(make):
    client, _ = make()
    assert "/static/export-image.js" in client.get("/").text
    js = client.get("/static/export-image.js").text
    assert "window.exportJpeg" in js and "image/jpeg" in js and "0.92" in js and "document.fonts.ready" in js


# ---- 微調：標題、重置、匯出資訊、金色 VFP ------------------------------------------


def test_result_title_and_reset_button_text(make):
    client, _ = make()
    page = client.get("/").text + client.get("/static/app.js").text
    assert "你的最終有價值產品(VFP)與副產品清單" in page
    assert "重置，開新案例" in page and "目前的對話會被清除，確定要重新開始嗎？" in page


def test_baton_is_gone_and_wand_is_used():
    assert not (WEB / "assets" / "icons" / "baton.svg").exists()
    for p in WEB.rglob("*"):
        if p.is_file() and p.suffix in {".html", ".css", ".js", ".svg"}:
            assert "baton" not in p.read_text(encoding="utf-8"), p.name


def test_info_line_industry_and_name():
    from datetime import date

    from coach.export import info_line

    d = date(2026, 10, 14)
    assert info_line({"industry": "餐飲業（餐廳、小吃、咖啡、飲料店）", "name": "林主管"}, d) == "產業：餐飲業　稱呼：林主管　日期：2026-10-14"
    assert info_line({"industry": "其他：修繕", "name": "王"}, d) == "產業：其他：修繕　稱呼：王　日期：2026-10-14"
    # 沒有稱呼／產業時整個欄位省略，不會出現 None 或 undefined
    for inputs in ({"industry": "資訊、軟體與網路服務", "name": ""}, {"industry": "資訊、軟體與網路服務"}):
        line = info_line(inputs, d)
        assert line == "產業：資訊、軟體與網路服務　日期：2026-10-14" and "None" not in line and "稱呼" not in line
    assert info_line(None, d) == "日期：2026-10-14"


def test_xlsx_info_row_has_industry_and_name(make):
    client, _ = make(vfp("V", ITEMS))
    sid = client.post("/sessions", json={"name": "林主管", "industry": 1}).json()["session_id"]
    ws = read_xlsx(client.get(f"/sessions/{sid}/export.xlsx"))["產品鑑定"]
    assert ws["A1"].value.startswith("產業：餐飲業　稱呼：林主管　日期：")
    assert "（" not in ws["A1"].value
    assert "由管到飛產生，日期：" in ws["A7"].value  # 最下方那行不變


def test_xlsx_info_row_for_industry_20(make):
    client, _ = make(vfp("V", ITEMS))
    sid = client.post("/sessions", json={"name": "林主管", "industry": 20, "industry_other": "修繕"}).json()["session_id"]
    ws = read_xlsx(client.get(f"/sessions/{sid}/export.xlsx"))["產品鑑定"]
    assert ws["A1"].value.startswith("產業：其他：修繕　稱呼：林主管")


def test_jpeg_script_has_info_line_gold_row_and_no_footer_brand(make):
    client, _ = make()
    js = client.get("/static/export-image.js").text
    assert "產業：" in js and "稱呼：" in js and "createLinearGradient" in js
    assert "#F4C84D" in js and "#FBE08A" in js
    assert 'fillText("管到飛"' not in js
    assert "管到飛｜最終有價值產品與副產品" in js  # 標題維持


def test_vfp_row_is_gold_in_css_without_white_text(make):
    client, _ = make()
    css = client.get("/static/style.css").text
    for color in ("#F4C84D", "#E0A92B", "#FFF3C9", "#FBE08A", "#D9A21B"):
        assert color in css
    assert "vfp-label { background: linear-gradient(135deg, #F4C84D, #E0A92B); color: var(--navy-900)" in css


def test_export_menu_opens_upward_aligned_right_and_actions_are_space_between(make):
    client, _ = make()
    css = client.get("/static/style.css").text
    assert "right: 0; bottom: calc(100% + 6px);" in css
    assert "justify-content: space-between; align-items: center; gap: 10px; }" in css
    html = client.get("/").text
    assert html.index('class="actions-left"') < html.index('id="export-btn"')


def test_home_background_is_plain_css_cover_right_top(make):
    client, _ = make()
    css = client.get("/static/style.css").text
    assert "background-size: cover; background-position: right top;" in css
    assert "filter: blur" not in css.split(".sky.has-photo")[1].split("}")[0]


def test_question_help_text_matches_rules_section_2_2(make):
    """說明與舉例的字必須同時出現在規則檔 2.2 與 app.js；規則檔改字而前端沒跟上就會失敗。"""
    rules = load_rules().text
    start = rules.index("### 2.2")
    section = rules[start : rules.index("###", start + 5)]
    js = make()[0].get("/static/app.js").text
    for phrase in ("寫越詳盡越好", "好吃的麵", "讓人吃了想再來買的麵包"):
        assert phrase in section, f"規則檔 2.2 沒有「{phrase}」"
        assert phrase in js, f"app.js 沒有「{phrase}」"


def test_question_help_is_text_only_and_removed_after_answering(make):
    client, _ = make()
    js = client.get("/static/app.js").text
    assert 'hint: "寫越詳盡越好"' in js and 'examples: ["好吃的麵", "讓人吃了想再來買的麵包"]' in js
    assert 'examples.join("　/　")' in js
    assert "clearQuestionHelp()" in js and ".q-help" in js
    css = client.get("/static/style.css").text
    assert ".q-help p { margin: 0; font-size: 14px;" in css and "color: var(--gray-600)" in css.split(".q-help p")[1].split("}")[0]


def test_six_questions_are_all_required_and_skip_is_gone(make):
    client, _ = make()
    html, js = client.get("/").text, client.get("/static/app.js").text
    intake_js = client.get("/static/intake.js").text  # 六題的狀態與驗證邏輯在 intake.js
    assert "略過" not in html and "略過" not in js and "skip" not in html.lower() and "skip" not in js.lower()
    assert "略過" not in intake_js and "skip" not in intake_js.lower()
    block = intake_js[intake_js.index("const FIELDS = [") : intake_js.index("const INDUSTRY_STEP")]
    assert block.count("required: true") == 6  # 六題都有必填設定
    assert block.count("key:") == 6
    # 逐題驗證：空白不送出、提醒顯示在輸入框下方、輸入框保留焦點
    assert "if (!t) return" in intake_js and 'EMPTY_HINT = "這題需要回答喔"' in intake_js
    assert 'id="answer-hint"' in html and "setAnswerHint(result.hint)" in js and '$("answer").focus()' in js


def test_score_must_be_1_to_10_and_industry_other_is_still_asked(make):
    client, _ = make()
    js = client.get("/static/app.js").text
    intake_js = client.get("/static/intake.js").text
    assert "n < 1 || n > 10" in intake_js and "分數請填 1～10 的整數喔" in intake_js
    assert 'OTHER_PROMPT = "請填寫你的產業"' in js and "isOther" in intake_js


def test_prev_button_is_wired_and_intake_script_loads_before_app(make):
    client, _ = make()
    html = client.get("/").text
    assert "← 上一題" in html and 'id="prev-btn"' in html
    assert html.index('id="prev-btn"') < html.index('id="answer-btn"')  # 在「送出」左側
    assert html.index("/static/intake.js") < html.index("/static/app.js")
    js = client.get("/static/app.js").text
    assert "Intake.back(intake)" in js and "Intake.canBack(intake)" in js
    assert '$("prev-btn").disabled = busy' in js  # 送出途中停用


# ---- 對話頁不露出 VFP；結果頁揭曉動畫 ------------------------------------------------


def _function_body(js, signature):
    start = js.index(signature)
    brace = js.index("{", start)
    depth = 0
    for i in range(brace, len(js)):
        depth += {"{": 1, "}": -1}.get(js[i], 0)
        if depth == 0:
            return js[start : i + 1]
    raise AssertionError(signature)


def test_chat_page_never_writes_vfp_into_the_dom(make):
    client, _ = make()
    js, html = client.get("/static/app.js").text, client.get("/").text
    # 產出之後的流程不接收 VFP，也沒有把 VFP 組成泡泡文字的函式
    assert "vfpMessage" not in js and "finishConversation(reply.vfp)" not in js
    finish = _function_body(js, "async function finishConversation(")
    assert "async function finishConversation()" in finish and "vfp" not in finish.lower()
    replay = _function_body(js, "function replaySession(")
    for forbidden in ("s.result", ".vfp", "byproducts", "edited", "original", "cur."):
        assert forbidden not in replay, forbidden
    assert replay.replace('s.status === "vfp"', "").lower().count("vfp") == 0
    assert 'const DONE_MESSAGE = "我整理好了，準備好看你的最終有價值產品了嗎？"' in js
    # 靜態頁面（對話頁）的文字、aria、title 裡也沒有結果內容
    chat_html = html[html.index('id="page-chat"') : html.index('id="page-result"')]
    assert "VFP" not in chat_html


def test_open_result_button_text_replaces_the_old_wording(make):
    client, _ = make()
    html, js = client.get("/").text, client.get("/static/app.js").text
    assert "打開最終結果" in html and "打開最終結果" in js
    for old in ("看副產品拆解", "前往結果"):
        assert old not in html and old not in js, old


def test_reveal_animation_only_from_button_and_respects_reduced_motion(make):
    client, _ = make()
    js, css = client.get("/static/app.js").text, client.get("/static/style.css").text
    assert "prefers-reduced-motion" in js and "prefers-reduced-motion" in css
    assert "revealNext = true" in js and "if (revealNext) startReveal()" in js
    # 重新整理、返回、編輯重畫都不播：flag 用完就清，離開結果頁也清，renderRows 會停掉動畫
    assert "revealNext = false;" in js and "stopReveal(); // 編輯造成重畫時不要重播" in js
    assert "600ms ease-out" in css and ".reveal #rows tr" in css
    assert 'id="result-card"' in client.get("/").text


def test_questions_config_keeps_original_text_and_names_are_added_at_display_time(make):
    client, _ = make()
    js = client.get("/static/app.js").text
    block = js[js.index("const QUESTIONS = [") : js.index("const OTHER_PROMPT")]
    assert "陳姐" not in block and "稱呼，" not in block  # 設定裡是不含稱呼的原文
    assert "Intake.displayQuestion(index, question, name)" in js
    # 所有顯示教練問句的地方都經過 askText；只有第 1 題以外才會加
    assert js.count("askText(") >= 5
    assert "addBubble(\"coach\", QUESTIONS[" not in js and "addBubble(\"coach\", q.question" not in js
    # 模型的追問不加稱呼
    assert 'addBubble("coach", reply.question)' in js and 'addBubble("coach", ex.coach.question)' in js
    # 驗證提醒不加稱呼
    intake_js = client.get("/static/intake.js").text
    assert 'EMPTY_HINT = "這題需要回答喔"' in intake_js


# ---- 舊格式（name / description）相容 ----------------------------------------------


def test_legacy_session_is_converted_on_read_without_rewriting_file(make, tmp_path):
    client, _ = make()
    legacy = {
        "id": "20250101-000000-abcdef",
        "status": "vfp",
        "pending": None,
        "inputs": {"name": "王", "industry": "餐飲業"},
        "exchanges": [],
        "result": {
            "created_at": "x", "rules_version": "v0", "rules_sha8": "00000000", "model": "m", "edited": None,
            "original": {"vfp": "舊VFP", "byproducts": [{"name": "備料", "description": "先備好"}]},
        },
    }
    d = tmp_path / "sessions"
    d.mkdir()
    path = d / f"{legacy['id']}.json"
    path.write_text(json.dumps(legacy, ensure_ascii=False), encoding="utf-8")
    before = path.read_text(encoding="utf-8")
    r = client.get(f"/sessions/{legacy['id']}")
    assert r.status_code == 200
    assert r.json()["result"]["original"]["byproducts"] == [{"task": "備料", "detail": "先備好", "product": "", "assignee": None}]
    x = client.get(f"/sessions/{legacy['id']}/export.xlsx")
    assert x.status_code == 200
    ws = read_xlsx(x)["產品鑑定"]
    assert [c.value for c in ws[4]] == [1, "備料", "先備好", None, "（未指派）"]
    assert path.read_text(encoding="utf-8") == before


def test_export_xlsx_product_column_wraps_and_aligns_top(make):
    client, _ = make(vfp("V", [{"task": "t", "detail": "d", "product": "很長的副產品" * 10}]))
    ws = read_xlsx(client.get(f"/sessions/{vfp_session(client)}/export.xlsx"))["產品鑑定"]
    assert ws["D4"].alignment.wrap_text and ws["D4"].alignment.vertical == "top"
    assert ws["D4"].value == "很長的副產品" * 10


def test_result_table_layout_is_fixed_with_action_column_width(make):
    client, _ = make()
    css = client.get("/static/style.css").text
    assert "table-layout: fixed" in css
    assert "--col-action: 124px" in css and ".c-act { width: var(--col-action); }" in css
    assert "--col-index: 56px" in css and "--col-assignee: 150px" in css
    assert "flex-wrap: nowrap" in css
    html = client.get("/").text
    assert '<div class="table-scroll">' in html and '<col class="c-act">' in html


def test_web_files_revalidate_and_index_has_version_params(make):
    client, _ = make()
    for path in ("/", "/static/app.js", "/static/style.css", "/static/export-image.js", "/static/index.html"):
        r = client.get(path)
        assert r.status_code == 200 and r.headers["cache-control"] == "no-cache", path
    assert client.get("/static/app.js").headers.get("etag")
    html = client.get("/").text
    for name in ("style.css", "export-image.js", "intake.js", "app.js"):
        assert re.search(rf'/static/{re.escape(name)}\?v=\d+"', html), name
    icons = Path(__file__).resolve().parent.parent / "coach" / "web" / "assets" / "icons"
    svg = next(icons.glob("*.svg"), None)
    if svg:
        assert client.get(f"/static/assets/icons/{svg.name}").headers.get("cache-control") != "no-cache"


def test_export_xlsx_tolerates_legacy_and_missing_values():
    from datetime import date

    from coach.export import build_xlsx

    wb = read_xlsx(type("R", (), {"content": build_xlsx("V", [{"name": "舊", "description": "d"}, {"task": "t", "detail": None, "product": None}, {}], date(2025, 1, 1))})())
    ws = wb["產品鑑定"]
    assert [c.value for c in ws[4]] == [1, "舊", "d", None, "（未指派）"]
    assert [c.value for c in ws[5]] == [2, "t", None, None, "（未指派）"]
    assert all("None" not in c.value and "undefined" not in c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str))


# ---- 員工名單與指派 ----------------------------------------------------------------


def put_staff(client, sid, staff):
    return client.put(f"/sessions/{sid}/staff", json={"staff": staff})


def test_staff_requires_vfp_state_409(make):
    client, _ = make(clarify("Q"))
    sid = client.post("/sessions", json=BODY).json()["session_id"]
    assert client.get(f"/sessions/{sid}/staff").status_code == 409
    assert put_staff(client, sid, [{"name": "甲"}]).status_code == 409


def test_staff_put_get_generates_ids_and_trims(make):
    client, _ = make(vfp("V", ITEMS))
    sid = vfp_session(client)
    r = put_staff(client, sid, [{"name": " 王小明 ", "role": " 店長 "}, {"name": "李小華"}])
    assert r.status_code == 200
    staff = r.json()["staff"]
    assert [(p["name"], p["role"]) for p in staff] == [("王小明", "店長"), ("李小華", "")]
    assert all(p["id"] and len(p["id"]) <= 8 for p in staff) and staff[0]["id"] != staff[1]["id"]
    assert client.get(f"/sessions/{sid}/staff").json() == {"staff": staff}
    assert client.get(f"/sessions/{sid}").json()["staff"] == staff
    # 改名後 id 不變
    r2 = put_staff(client, sid, [{"id": staff[0]["id"], "name": "王大明", "role": ""}, staff[1]])
    assert r2.json()["staff"][0]["id"] == staff[0]["id"] and r2.json()["staff"][0]["name"] == "王大明"


@pytest.mark.parametrize(
    "staff,errors",
    [
        ([{"name": "甲"}, {"name": "  "}], [{"row": 2, "field": "name", "message": "姓名不得為空"}]),
        ([{"name": "甲"}, {"name": "甲 "}], [{"row": 2, "field": "name", "message": "這個姓名已經在名單中"}]),
        ([{"name": "長" * 21}], [{"row": 1, "field": "name", "message": "姓名最多 20 字"}]),
        ([{"name": f"人{i}"} for i in range(51)], [{"row": 51, "field": "name", "message": "員工最多 50 人"}]),
    ],
)
def test_staff_validation_422_with_errors(make, staff, errors):
    client, _ = make(vfp("V", ITEMS))
    sid = vfp_session(client)
    r = put_staff(client, sid, staff)
    assert r.status_code == 422 and isinstance(r.json()["detail"], str)
    assert r.json()["errors"] == errors
    assert client.get(f"/sessions/{sid}").json()["staff"] == []  # 失敗不存


def test_staff_limit_50_ok(make):
    client, _ = make(vfp("V", ITEMS))
    sid = vfp_session(client)
    assert put_staff(client, sid, [{"name": f"人{i}"} for i in range(50)]).status_code == 200


def test_assignee_must_exist_and_deleting_staff_unassigns(make):
    client, _ = make(vfp("V", ITEMS))
    sid = vfp_session(client)
    a, b = put_staff(client, sid, [{"name": "甲"}, {"name": "乙"}]).json()["staff"]
    items = [{"task": "t1", "detail": "", "product": "p1", "assignee": a["id"]}, {"task": "t2", "detail": "", "product": "p2", "assignee": b["id"]}]
    r = client.put(f"/sessions/{sid}/byproducts", json={"vfp": "V", "byproducts": items})
    assert r.status_code == 200 and [x["assignee"] for x in r.json()["byproducts"]] == [a["id"], b["id"]]
    bad = [items[0], {**items[1], "assignee": "nope"}]
    r = client.put(f"/sessions/{sid}/byproducts", json={"vfp": "V", "byproducts": bad})
    assert r.status_code == 422
    assert r.json()["errors"] == [{"row": 2, "field": "assignee", "message": "負責人不在員工名單中"}]
    assert "第 2 列" in r.json()["detail"]
    # 刪掉甲：原本指派給甲的項目改成未指派，乙的不受影響
    put_staff(client, sid, [b])
    cur = client.get(f"/sessions/{sid}").json()["result"]["edited"]["byproducts"]
    assert [x["assignee"] for x in cur] == [None, b["id"]]


def test_byproducts_422_errors_list_rows_and_fields(make):
    client, _ = make(vfp("V", ITEMS))
    sid = vfp_session(client)
    body = {"vfp": "V", "byproducts": [{"task": "a", "product": "p"}, {"task": " ", "product": ""}, {"task": "c", "product": "p", "detail": ""}]}
    r = client.put(f"/sessions/{sid}/byproducts", json=body)
    assert r.status_code == 422
    assert r.json()["errors"] == [
        {"row": 2, "field": "task", "message": "工作項目不得為空"},
        {"row": 2, "field": "product", "message": "副產品不得為空"},
    ]
    assert r.json()["detail"] == "第 2 列的工作項目不可為空"


def test_legacy_session_without_staff_or_assignee_gets_defaults(make, tmp_path):
    client, _ = make()
    legacy = {
        "id": "20250101-000001-abcdef", "status": "vfp", "pending": None, "inputs": {"name": "王", "industry": "餐飲業"}, "exchanges": [],
        "result": {"created_at": "x", "rules_version": "v0", "rules_sha8": "00000000", "model": "m", "edited": None,
                   "original": {"vfp": "V", "byproducts": [{"task": "t", "detail": "d", "product": "p"}]}},
    }
    d = tmp_path / "sessions"
    d.mkdir()
    path = d / f"{legacy['id']}.json"
    path.write_text(json.dumps(legacy, ensure_ascii=False), encoding="utf-8")
    before = path.read_text(encoding="utf-8")
    body = client.get(f"/sessions/{legacy['id']}").json()
    assert body["staff"] == [] and body["result"]["original"]["byproducts"][0]["assignee"] is None
    assert client.get(f"/sessions/{legacy['id']}/staff").json() == {"staff": []}
    assert path.read_text(encoding="utf-8") == before


def test_export_xlsx_has_assignee_column_and_by_assignee_sheet(make):
    client, _ = make(vfp("V", ITEMS))
    sid = vfp_session(client)
    a, b = put_staff(client, sid, [{"name": "王小明", "role": "店長"}, {"name": "李小華"}]).json()["staff"]
    items = [
        {"task": "t1", "detail": "d1", "product": "p1", "assignee": b["id"]},
        {"task": "t2", "detail": "", "product": "p2", "assignee": None},
        {"task": "t3", "detail": "", "product": "p3", "assignee": a["id"]},
        {"task": "t4", "detail": "", "product": "p4", "assignee": b["id"]},
    ]
    assert client.put(f"/sessions/{sid}/byproducts", json={"vfp": "V", "byproducts": items}).status_code == 200
    wb = read_xlsx(client.get(f"/sessions/{sid}/export.xlsx"))
    ws = wb["產品鑑定"]
    assert [c.value for c in ws[3]] == ["序號", "工作項目", "工作內容", "副產品", "負責人"]
    assert [ws.cell(row=r, column=5).value for r in (4, 5, 6, 7)] == ["李小華", "（未指派）", "王小明（店長）", "李小華"]
    assert any(str(r) in (ws.merged_cells.ranges.__str__()) for r in ("A1:E1", "B2:E2"))
    assert "A1:E1" in {str(r) for r in ws.merged_cells.ranges} and "B2:E2" in {str(r) for r in ws.merged_cells.ranges}
    assert "由管到飛產生" in ws["A9"].value
    by = wb["依負責人"]
    assert [c.value for c in by[1]] == ["負責人", "序號", "工作項目", "副產品"]
    # 依名單順序：王小明、李小華，未指派最後
    rows = [[c.value for c in by[i]] for i in range(2, 6)]
    assert [(r[0] or "", r[2]) for r in rows] == [("王小明（店長）", "t3"), ("李小華", "t1"), ("", "t4"), ("（未指派）", "t2")]


def test_result_page_has_assignee_column_and_staff_dialog_served_locally(make):
    client, _ = make()
    html = client.get("/").text
    assert '<col class="c-assignee">' in html and "負責人" in html
    assert 'id="staff-dialog"' in html and 'id="form-errors"' in html and "員工名單（0 人）" in html
    assert re.search(r'/static/staff\.js\?v=\d+"', html)
    assert client.get("/static/staff.js").status_code == 200
    assert "http://" not in html.replace("http://www.w3.org", "") and "https://" not in html
    css = client.get("/static/style.css").text
    assert ".c-assignee { width: var(--col-assignee); }" in css and ".field-error" in css and "#B3261E" in css


def test_voice_input_has_mic_button_and_no_external_urls(make):
    client, _ = make()
    html = client.get("/").text
    assert 'id="mic-btn"' in html and 'aria-label="語音輸入"' in html and "語音由瀏覽器的辨識服務處理" in html
    assert re.search(r'/static/voice\.js\?v=\d+"', html)
    for name in ("voice.js", "app.js", "intake.js"):
        js = client.get(f"/static/{name}").text
        assert not re.search(r"https?://", js), name
    app_js = client.get("/static/app.js").text
    for needle in ('"zh-TW"', "interimResults = true", "continuous = false", "maxAlternatives = 1"):
        assert needle in app_js
